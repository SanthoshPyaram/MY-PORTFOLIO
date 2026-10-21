import os
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_report():
    doc = docx.Document()
    
    # Configure 0.75-inch page margins for clean presentation
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
        # Header: Subtle watermark-style student identification
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("24MIC7312  |  SWE3004 FRONTEND DESIGN AND TESTING")
        hrun.font.name = "Arial"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(148, 163, 184)
        
        # Footer: Institutional details
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        frun1 = fp.add_run("VIT-AP UNIVERSITY  •  School of Computer Science & Engineering")
        frun1.font.name = "Arial"
        frun1.font.size = Pt(8.5)
        frun1.font.color.rgb = RGBColor(148, 163, 184)

    # Helper: Set cell background shading
    def set_cell_background(cell, hex_color):
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        cell._tc.get_or_add_tcPr().append(shd)

    # Helper: Add code block
    def add_code_section(doc, title, filepath, code_text):
        # Section Heading
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(2)
        h.paragraph_format.keep_with_next = True
        hrun = h.add_run(title)
        hrun.bold = True
        hrun.font.name = "Arial"
        hrun.font.size = Pt(13)
        hrun.font.color.rgb = RGBColor(15, 23, 42)

        # Filepath line
        fp = doc.add_paragraph()
        fp.paragraph_format.space_before = Pt(0)
        fp.paragraph_format.space_after = Pt(2)
        fp.paragraph_format.keep_with_next = True
        fprun = fp.add_run(filepath)
        fprun.bold = True
        fprun.font.name = "Consolas"
        fprun.font.size = Pt(10)
        fprun.font.color.rgb = RGBColor(37, 99, 235)

        # Label CODE
        cp = doc.add_paragraph()
        cp.paragraph_format.space_before = Pt(0)
        cp.paragraph_format.space_after = Pt(4)
        cp.paragraph_format.keep_with_next = True
        crun = cp.add_run("CODE")
        crun.bold = True
        crun.font.name = "Arial"
        crun.font.size = Pt(9.5)
        crun.font.color.rgb = RGBColor(71, 85, 105)

        # Monospace code chunked for Word performance
        lines = code_text.splitlines()
        chunk_size = 50
        for i in range(0, len(lines), chunk_size):
            chunk = "\n".join(lines[i:i + chunk_size])
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(chunk)
            run.font.name = "Consolas"
            run.font.size = Pt(8.0)
            run.font.color.rgb = RGBColor(30, 41, 59)

    # Helper: Add Output/Screenshot section
    def add_screenshot_section(doc, label, image_path, caption=None, width=Inches(6.2)):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(f"OUTPUT / SCREENSHOT — {label}")
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(15, 23, 42)

        if os.path.exists(image_path):
            img_p = doc.add_paragraph()
            img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            img_p.paragraph_format.space_before = Pt(4)
            img_p.paragraph_format.space_after = Pt(4)
            img_run = img_p.add_run()
            img_run.add_picture(image_path, width=width)
            
            if caption:
                cap_p = doc.add_paragraph()
                cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap_p.paragraph_format.space_before = Pt(0)
                cap_p.paragraph_format.space_after = Pt(10)
                cap_run = cap_p.add_run(caption)
                cap_run.font.name = "Arial"
                cap_run.font.size = Pt(8.5)
                cap_run.font.italic = True
                cap_run.font.color.rgb = RGBColor(100, 116, 139)

    # =========================================================================
    # PAGE 1: TITLE & COVER DETAILS
    # =========================================================================
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(18)
    title_p.paragraph_format.space_after = Pt(4)
    run_t1 = title_p.add_run("FRONTEND DESIGN AND TESTING LAB\n")
    run_t1.bold = True
    run_t1.font.name = "Arial"
    run_t1.font.size = Pt(16)
    run_t1.font.color.rgb = RGBColor(30, 41, 59)

    run_t2 = title_p.add_run("DEVELOP A PERSONAL PORTFOLIO WEBSITE (ANGULAR & TYPESCRIPT)")
    run_t2.bold = True
    run_t2.font.name = "Arial"
    run_t2.font.size = Pt(15)
    run_t2.font.color.rgb = RGBColor(30, 41, 59)

    # Google Drive / GitHub Link
    link_p = doc.add_paragraph()
    link_p.paragraph_format.space_before = Pt(20)
    link_p.paragraph_format.space_after = Pt(2)
    l_label = link_p.add_run("GOOGLE DRIVE / GITHUB REPOSITORY LINK:\n")
    l_label.bold = True
    l_label.font.name = "Arial"
    l_label.font.size = Pt(11)
    l_label.font.color.rgb = RGBColor(67, 56, 202)

    l_val = link_p.add_run("https://github.com/SanthoshPyaram/MY-PORTFOLIO.git")
    l_val.font.name = "Arial"
    l_val.font.size = Pt(10.5)
    l_val.font.underline = True
    l_val.font.color.rgb = RGBColor(37, 99, 235)

    # Files Included Header
    fi_p = doc.add_paragraph()
    fi_p.paragraph_format.space_before = Pt(22)
    fi_p.paragraph_format.space_after = Pt(2)
    fi_run = fi_p.add_run("Files included (Angular & TypeScript Architecture)")
    fi_run.bold = True
    fi_run.font.name = "Arial"
    fi_run.font.size = Pt(12)
    fi_run.font.color.rgb = RGBColor(15, 23, 42)

    count_p = doc.add_paragraph()
    count_p.paragraph_format.space_before = Pt(0)
    count_p.paragraph_format.space_after = Pt(10)
    c_run = count_p.add_run("7 core Angular & TypeScript production source modules are included in this complete assignment submission:")
    c_run.font.name = "Arial"
    c_run.font.size = Pt(9.5)
    c_run.font.color.rgb = RGBColor(71, 85, 105)

    # List of Files
    file_list = [
        ("src/app/models/portfolio.models.ts", "TypeScript interfaces for Projects, Skills, Education, Certifications, and Social Profiles"),
        ("src/app/services/portfolio.service.ts", "Angular Injectable state service with 4 featured projects (including 2 ECS projects) and credentials"),
        ("src/app/services/theme.service.ts", "Angular reactive theme switching service with localStorage persistence and signal-based state"),
        ("src/app/components/projects/projects.component.ts", "Angular component rendering 4 project cards with detail modal trigger bindings"),
        ("src/app/components/skills/skills.component.ts", "Angular component rendering skill matrix, progress bars, and Embedded Systems & IoT (ECS) category"),
        ("src/app/components/resume-modal/resume-modal.component.ts", "Angular component rendering printable resume document with print-to-PDF handler"),
        ("src/styles.css", "Design tokens, Charcoal & Cobalt palette, glassmorphism, responsive grid, animations, and print media queries")
    ]

    for fname, fdesc in file_list:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_before = Pt(0)
        bp.paragraph_format.space_after = Pt(3)
        b_name = bp.add_run(f"{fname} ")
        b_name.bold = True
        b_name.font.name = "Consolas"
        b_name.font.size = Pt(9.5)
        b_name.font.color.rgb = RGBColor(37, 99, 235)

        b_desc = bp.add_run(f"— {fdesc}")
        b_desc.font.name = "Arial"
        b_desc.font.size = Pt(9.0)
        b_desc.font.color.rgb = RGBColor(71, 85, 105)

    # Page Break before code modules
    doc.add_page_break()

    # =========================================================================
    # MODULE 1: Data Models & TypeScript Interfaces
    # =========================================================================
    with open("src/app/models/portfolio.models.ts", "r", encoding="utf-8") as f:
        models_code = f.read()

    add_code_section(
        doc,
        title="1. TypeScript Data Models & Contract Definitions",
        filepath="src/app/models/portfolio.models.ts",
        code_text=models_code
    )

    add_screenshot_section(
        doc,
        label="Home / Hero Section & Sticky Navigation Bar",
        image_path=r"screenshots\1_home_hero.png",
        caption="Figure 1.1: Angular application Home hero section displaying brand navbar, developer headline, tech halo avatar, and primary action buttons.",
        width=Inches(6.2)
    )

    doc.add_page_break()

    # =========================================================================
    # MODULE 2: Portfolio Data & State Service
    # =========================================================================
    with open("src/app/services/portfolio.service.ts", "r", encoding="utf-8") as f:
        service_code = f.read()

    add_code_section(
        doc,
        title="2. Angular Portfolio State & Data Service",
        filepath="src/app/services/portfolio.service.ts",
        code_text=service_code
    )

    add_screenshot_section(
        doc,
        label="About Me Section & Academic Credentials",
        image_path=r"screenshots\2_about_section.png",
        caption="Figure 2.1: Angular application About section showcasing personal narrative, verified academic scores (9.68 CGPA, 982/1000 Intermediate, 589/600 SSC), and profile portrait.",
        width=Inches(6.2)
    )

    doc.add_page_break()

    # =========================================================================
    # MODULE 3: Reactive Theme Service
    # =========================================================================
    with open("src/app/services/theme.service.ts", "r", encoding="utf-8") as f:
        theme_code = f.read()

    add_code_section(
        doc,
        title="3. Angular Reactive Theme Switching Service",
        filepath="src/app/services/theme.service.ts",
        code_text=theme_code
    )

    add_screenshot_section(
        doc,
        label="Technical Skills Section & Embedded ECS Category",
        image_path=r"screenshots\3_skills_section.png",
        caption="Figure 3.1: Angular application Skills section demonstrating responsive category cards, animated progress bars, and the Embedded Systems & IoT (ECS) skill matrix.",
        width=Inches(6.2)
    )

    doc.add_page_break()

    # =========================================================================
    # MODULE 4: Projects Component
    # =========================================================================
    with open("src/app/components/projects/projects.component.ts", "r", encoding="utf-8") as f:
        projects_comp_code = f.read()

    add_code_section(
        doc,
        title="4. Angular Projects Component (4 Featured & ECS Projects)",
        filepath="src/app/components/projects/projects.component.ts",
        code_text=projects_comp_code
    )

    add_screenshot_section(
        doc,
        label="Featured Projects Showcase (4 Projects)",
        image_path=r"screenshots\4_projects_section.png",
        caption="Figure 4.1: Angular application Projects showcase featuring Smart Agriculture, COOPNEX, Smart Shoes for Blind Persons, and Smart Air Purification System.",
        width=Inches(6.2)
    )

    doc.add_page_break()

    # =========================================================================
    # MODULE 5: Skills Component
    # =========================================================================
    with open("src/app/components/skills/skills.component.ts", "r", encoding="utf-8") as f:
        skills_comp_code = f.read()

    add_code_section(
        doc,
        title="5. Angular Skills Component & Progress Indicators",
        filepath="src/app/components/skills/skills.component.ts",
        code_text=skills_comp_code
    )

    add_screenshot_section(
        doc,
        label="Interactive Contact Section with TypeScript Form",
        image_path=r"screenshots\5_contact_section.png",
        caption="Figure 5.1: Angular application Contact section featuring interactive message form with real-time field validation, contact channels, and social media connectivity.",
        width=Inches(6.2)
    )

    doc.add_page_break()

    # =========================================================================
    # MODULE 6: Resume Modal Component
    # =========================================================================
    with open("src/app/components/resume-modal/resume-modal.component.ts", "r", encoding="utf-8") as f:
        resume_comp_code = f.read()

    add_code_section(
        doc,
        title="6. Angular Resume Modal Component with Print Handler",
        filepath="src/app/components/resume-modal/resume-modal.component.ts",
        code_text=resume_comp_code
    )

    add_screenshot_section(
        doc,
        label="Dynamic Interactive Resume Modal View",
        image_path=r"screenshots\6_resume_modal.png",
        caption="Figure 6.1: Angular application Resume modal rendered dynamically with print-to-PDF formatting, exact 9.68 CGPA, and all 4 projects.",
        width=Inches(6.2)
    )

    doc.add_page_break()

    # =========================================================================
    # MODULE 7: Stylesheet & Design System
    # =========================================================================
    with open("src/styles.css", "r", encoding="utf-8") as f:
        styles_code = f.read()

    add_code_section(
        doc,
        title="7. Global Styles, Design Tokens & Glassmorphism System",
        filepath="src/styles.css",
        code_text=styles_code
    )

    doc.add_page_break()

    # =========================================================================
    # FINAL PAGE: EXACT ACADEMIC SUBMISSION DETAILS TABLE
    # =========================================================================
    final_p = doc.add_paragraph()
    final_p.paragraph_format.space_before = Pt(36)
    final_p.paragraph_format.space_after = Pt(16)
    final_run = final_p.add_run("LAB SUBMISSION & VERIFICATION DETAILS")
    final_run.bold = True
    final_run.font.name = "Arial"
    final_run.font.size = Pt(14)
    final_run.font.color.rgb = RGBColor(15, 23, 42)

    # 8-Row x 2-Column Gold / Yellow Academic Table (Exact replica of Reference Page 32)
    table = doc.add_table(rows=8, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    for row in table.rows:
        row.cells[0].width = Inches(2.2)
        row.cells[1].width = Inches(4.3)

    table_data = [
        ("REGNO", "24MIC7312"),
        ("NAME", "P SANTHOSH REDDY"),
        ("COURSE", "FRONTEND DESIGN AND TESTING"),
        ("COURSE CODE", "SWE3004"),
        ("SLOT", "L41+L42"),
        ("DATE", "21-october-2026"),
        ("FACULTY", "KANDRU LAKSHMI SAI PRANEETH"),
        ("VENUE", "CB-409")
    ]

    for idx, (label, val) in enumerate(table_data):
        row = table.rows[idx]
        
        # Cell 1 (Label)
        c1 = row.cells[0]
        c1.text = label
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(4)
        p1.paragraph_format.space_after = Pt(4)
        r1 = p1.runs[0]
        r1.bold = True
        r1.font.name = "Arial"
        r1.font.size = Pt(11)
        r1.font.color.rgb = RGBColor(0, 0, 0)
        set_cell_background(c1, "FDC02F")
        
        # Cell 2 (Value)
        c2 = row.cells[1]
        c2.text = val
        p2 = c2.paragraphs[0]
        p2.paragraph_format.space_before = Pt(4)
        p2.paragraph_format.space_after = Pt(4)
        r2 = p2.runs[0]
        r2.bold = True
        r2.font.name = "Arial"
        r2.font.size = Pt(11)
        r2.font.color.rgb = RGBColor(0, 0, 0)
        set_cell_background(c2, "FDC02F")

    # Set solid black borders on the table
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="12" w:space="0" w:color="000000"/>'
        f'<w:bottom w:val="single" w:sz="12" w:space="0" w:color="000000"/>'
        f'<w:left w:val="single" w:sz="12" w:space="0" w:color="000000"/>'
        f'<w:right w:val="single" w:sz="12" w:space="0" w:color="000000"/>'
        f'<w:insideH w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
        f'<w:insideV w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

    # Repository Link below table
    gl_p = doc.add_paragraph()
    gl_p.paragraph_format.space_before = Pt(36)
    gl_p.paragraph_format.space_after = Pt(2)
    gl_lbl = gl_p.add_run("GOOGLE DRIVE / GITHUB LINK:\n")
    gl_lbl.bold = True
    gl_lbl.font.name = "Arial"
    gl_lbl.font.size = Pt(10.5)
    gl_lbl.font.color.rgb = RGBColor(67, 56, 202)

    gl_val = gl_p.add_run("https://github.com/SanthoshPyaram/MY-PORTFOLIO.git")
    gl_val.font.name = "Arial"
    gl_val.font.size = Pt(10)
    gl_val.font.underline = True
    gl_val.font.color.rgb = RGBColor(37, 99, 235)

    # Save document locally
    out_filename = "FRONTEND_DESIGN_AND_TESTING_LAB_PORTFOLIO_24MIC7312.docx"
    doc.save(out_filename)
    print(f"Document successfully created and saved as: {out_filename}")
    print(f"File size: {os.path.getsize(out_filename)} bytes")

    # Automatically copy to user Downloads and Desktop directories
    dest_downloads = os.path.join(r"C:\Users\SANTHOSH PYARAM\Downloads", out_filename)
    dest_desktop = os.path.join(r"C:\Users\SANTHOSH PYARAM\Desktop", out_filename)
    dest_myfiles_desktop = os.path.join(r"C:\Users\SANTHOSH PYARAM\Downloads\MY FILES\Desktop", out_filename)

    for dest in [dest_downloads, dest_desktop, dest_myfiles_desktop]:
        try:
            shutil.copy2(out_filename, dest)
            print(f"Copied DOCX to: {dest}")
        except Exception as e:
            print(f"Copy error to {dest}: {e}")

if __name__ == "__main__":
    create_report()
